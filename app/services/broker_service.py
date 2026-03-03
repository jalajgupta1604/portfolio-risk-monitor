import csv
import logging
import uuid
from io import BytesIO, StringIO

from fastapi import HTTPException, status
from sqlalchemy.ext.asyncio import AsyncSession

from app.brokers import BrokerHolding, BrokerRegistry, decrypt_token, encrypt_token
from app.models import Holding, Portfolio
from app.repositories.broker_repo import BrokerConnectionRepository
from app.repositories.portfolio_repo import PortfolioRepository
from app.services.stock_service import StockService
from app.schemas.broker import (
    AutoSyncResponse,
    BrokerConnectionResponse,
    BrokerConnectionsListResponse,
    BrokerInfoResponse,
    BrokerListResponse,
    CsvImportResponse,
    ImportedHoldingSummary,
    ImportMode,
    SyncResponse,
    ZerodhaLoginUrlResponse,
)

logger = logging.getLogger(__name__)


class BrokerService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.broker_repo = BrokerConnectionRepository(session)
        self.portfolio_repo = PortfolioRepository(session)
        self.registry = BrokerRegistry()

    # ── Available brokers ──

    def list_available_brokers(self) -> BrokerListResponse:
        brokers = self.registry.list_all()
        return BrokerListResponse(
            brokers=[
                BrokerInfoResponse(
                    name=b.name,
                    display_name=b.display_name,
                    broker_type=b.broker_type,
                    supports_api=b.supports_api,
                    supports_csv=b.supports_csv,
                )
                for b in brokers
            ]
        )

    # ── Connections management ──

    async def list_connections(
        self, user_id: uuid.UUID
    ) -> BrokerConnectionsListResponse:
        conns = await self.broker_repo.list_by_user(user_id)
        return BrokerConnectionsListResponse(
            connections=[
                BrokerConnectionResponse(
                    id=c.id,
                    broker_name=c.broker_name,
                    is_active=c.is_active,
                    last_synced_at=c.last_synced_at,
                    created_at=c.created_at,
                )
                for c in conns
            ]
        )

    async def delete_connection(
        self, connection_id: uuid.UUID, user_id: uuid.UUID
    ) -> None:
        conn = await self.broker_repo.get_by_id(connection_id, user_id=user_id)
        if not conn:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Broker connection not found",
            )
        await self.broker_repo.delete(connection_id)

    # ── Zerodha OAuth ──

    def get_zerodha_login_url(self) -> ZerodhaLoginUrlResponse:
        broker = self.registry.get_api_broker("zerodha")
        if not broker:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Zerodha API integration is not configured",
            )
        return ZerodhaLoginUrlResponse(login_url=broker.get_login_url())

    async def handle_zerodha_callback(
        self, request_token: str, user_id: uuid.UUID
    ) -> BrokerConnectionResponse:
        broker = self.registry.get_api_broker("zerodha")
        if not broker:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Zerodha API integration is not configured",
            )
        try:
            token_data = broker.exchange_token(request_token)
        except Exception as e:
            logger.error("Zerodha token exchange failed: %s", e)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Zerodha token exchange failed: {e}",
            )

        encrypted = encrypt_token(token_data["access_token"])

        # Upsert: update existing connection or create new
        existing = await self.broker_repo.get_by_user_and_broker(user_id, "zerodha")
        if existing:
            conn = await self.broker_repo.update_token(existing, encrypted)
        else:
            conn = await self.broker_repo.create(
                user_id=user_id,
                broker_name="zerodha",
                access_token_encrypted=encrypted,
                api_key=broker.api_key,
            )

        return BrokerConnectionResponse(
            id=conn.id,
            broker_name=conn.broker_name,
            is_active=conn.is_active,
            last_synced_at=conn.last_synced_at,
            created_at=conn.created_at,
        )

    # ── File import (CSV / PDF) ──

    async def import_holdings_file(
        self,
        user_id: uuid.UUID,
        broker_name: str,
        import_mode: ImportMode,
        raw_content: bytes,
        portfolio_id: uuid.UUID | None = None,
        portfolio_name: str | None = None,
    ) -> CsvImportResponse:
        # Try PDF parser first, then CSV parser
        pdf_parser = self.registry.get_pdf_parser(broker_name)
        csv_parser = self.registry.get_csv_parser(broker_name)

        if not pdf_parser and not csv_parser:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unknown broker: {broker_name}",
            )

        try:
            if pdf_parser:
                broker_holdings = pdf_parser.parse_pdf(raw_content)
            else:
                csv_text = self._read_as_csv(raw_content)
                broker_holdings = csv_parser.parse(StringIO(csv_text))
        except HTTPException:
            raise
        except ValueError as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Parse error: {e}",
            )

        if not broker_holdings:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="No valid holdings found in file",
            )

        if import_mode == ImportMode.MERGE_INTO:
            if not portfolio_id:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="portfolio_id is required for merge_into mode",
                )
            portfolio = await self.portfolio_repo.get_by_id(portfolio_id, user_id=user_id)
            if not portfolio:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail="Portfolio not found",
                )
        else:
            active_parser = pdf_parser or csv_parser
            name = portfolio_name or f"Import from {active_parser.display_name}"
            portfolio = await self.portfolio_repo.create(
                name=name, user_id=user_id
            )
            # Eagerly load holdings to avoid lazy-load in async context
            await self.session.refresh(portfolio, ["holdings"])

        summaries = await self._apply_holdings(
            portfolio, broker_holdings, source=broker_name
        )

        return CsvImportResponse(
            portfolio_id=portfolio.id,
            portfolio_name=portfolio.name,
            total_imported=len(summaries),
            created=sum(1 for s in summaries if s.action == "created"),
            updated=sum(1 for s in summaries if s.action == "updated"),
            holdings=summaries,
        )

    # ── Sync (API broker) ──

    async def sync_connection(
        self, connection_id: uuid.UUID, user_id: uuid.UUID
    ) -> SyncResponse:
        conn = await self.broker_repo.get_by_id(connection_id, user_id=user_id)
        if not conn:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Broker connection not found",
            )
        if not conn.is_active or not conn.access_token_encrypted:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Broker connection is not active or has no token",
            )

        broker = self.registry.get_api_broker(conn.broker_name)
        if not broker:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Broker {conn.broker_name} not available",
            )

        try:
            access_token = decrypt_token(conn.access_token_encrypted)
            broker_holdings = broker.fetch_holdings(access_token)
        except Exception as e:
            logger.error("Sync failed for %s: %s", conn.broker_name, e)
            await self.broker_repo.deactivate(conn.id)
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Sync failed (token may have expired): {e}",
            )

        # Find or create a portfolio for this broker connection
        portfolios = await self.portfolio_repo.list_all(user_id=user_id)
        # Use first portfolio or create one for the broker
        portfolio = next(
            (p for p in portfolios if any(
                h.source == conn.broker_name for h in p.holdings
            )),
            None,
        )
        if not portfolio:
            if portfolios:
                portfolio = portfolios[0]
            else:
                portfolio = await self.portfolio_repo.create(
                    name=f"{broker.display_name} Holdings",
                    user_id=user_id,
                )

        summaries = await self._apply_holdings(
            portfolio, broker_holdings, source=conn.broker_name
        )
        await self.broker_repo.update_last_synced(conn.id)

        return SyncResponse(
            broker_name=conn.broker_name,
            holdings_synced=len(summaries),
            portfolio_id=portfolio.id,
            created=sum(1 for s in summaries if s.action == "created"),
            updated=sum(1 for s in summaries if s.action == "updated"),
        )

    async def auto_sync_all(self, user_id: uuid.UUID) -> AutoSyncResponse:
        conns = await self.broker_repo.list_active_api_connections(user_id)
        synced: list[SyncResponse] = []
        errors: list[str] = []

        for conn in conns:
            try:
                result = await self.sync_connection(conn.id, user_id)
                synced.append(result)
            except Exception as e:
                logger.warning("Auto-sync failed for %s: %s", conn.broker_name, e)
                errors.append(f"{conn.broker_name}: {e}")

        return AutoSyncResponse(synced=synced, errors=errors)

    # ── Helpers ──

    @staticmethod
    def _read_as_csv(raw_content: bytes) -> str:
        """Convert raw file bytes to CSV text. Handles .csv and .xlsx.

        For xlsx files (e.g. Zerodha Console exports), skips leading
        metadata/summary rows and finds the actual data table by looking
        for the first row with 3+ non-empty cells.
        """
        # XLSX magic bytes: PK (zip archive)
        if raw_content[:4] == b"PK\x03\x04":
            try:
                import openpyxl

                wb = openpyxl.load_workbook(BytesIO(raw_content))
                # Prefer "Equity" sheet if present, else use active sheet
                ws = wb["Equity"] if "Equity" in wb.sheetnames else wb.active
                out = StringIO()
                writer = csv.writer(out)

                # Skip metadata rows: find first row with 3+ non-empty cells
                found_header = False
                for row in ws.iter_rows(values_only=True):
                    non_empty = [c for c in row if c is not None and str(c).strip()]
                    if not found_header:
                        if len(non_empty) >= 3:
                            found_header = True
                        else:
                            continue
                    # Once we found the header, skip fully blank rows
                    if not non_empty:
                        continue
                    writer.writerow(
                        [str(cell).strip() if cell is not None else "" for cell in row]
                    )
                wb.close()

                if not found_header:
                    raise ValueError("No data table found in the spreadsheet")
                return out.getvalue()
            except ImportError:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="XLSX support requires openpyxl. Please install it or convert to CSV first.",
                )
            except ValueError as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Failed to read XLSX file: {e}",
                )
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Failed to read XLSX file: {e}",
                )

        # Otherwise treat as CSV text
        try:
            return raw_content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Unsupported file format. Please upload a .csv or .xlsx file.",
            )

    async def _apply_holdings(
        self,
        portfolio: Portfolio,
        broker_holdings: list[BrokerHolding],
        source: str,
    ) -> list[ImportedHoldingSummary]:
        """Upsert holdings into portfolio: update existing symbols, add new ones."""
        existing_map: dict[str, Holding] = {
            h.symbol: h for h in portfolio.holdings
        }
        summaries: list[ImportedHoldingSummary] = []

        for bh in broker_holdings:
            # Resolve sector: prefer broker data, fallback to yfinance
            sector = bh.sector
            if not sector:
                try:
                    sector = await StockService.get_sector(bh.symbol)
                except Exception:
                    logger.warning("Failed to resolve sector for %s", bh.symbol)

            if bh.symbol in existing_map:
                # Update existing holding
                holding = existing_map[bh.symbol]
                holding.quantity = bh.quantity
                holding.avg_buy_price = bh.avg_buy_price
                if bh.current_price > 0:
                    holding.current_price = bh.current_price
                holding.source = source
                if sector:
                    holding.sector = sector
                summaries.append(
                    ImportedHoldingSummary(
                        symbol=bh.symbol,
                        quantity=bh.quantity,
                        avg_buy_price=bh.avg_buy_price,
                        action="updated",
                    )
                )
            else:
                # Create new holding
                holding = await self.portfolio_repo.add_holding(
                    portfolio_id=portfolio.id,
                    symbol=bh.symbol,
                    quantity=bh.quantity,
                    avg_buy_price=bh.avg_buy_price,
                    current_price=bh.current_price,
                    sector=sector,
                )
                holding.source = source
                summaries.append(
                    ImportedHoldingSummary(
                        symbol=bh.symbol,
                        quantity=bh.quantity,
                        avg_buy_price=bh.avg_buy_price,
                        action="created",
                    )
                )

        await self.session.flush()
        return summaries
